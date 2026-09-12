import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import pandas as pd
import openpyxl
from openpyxl.styles import Border, Side, Alignment, Font, PatternFill
from datetime import datetime
import os
import threading
import tempfile
import json

# ==============================================================================
# SAFE IMPORT: Mencegah crash jika modul sync_master belum tersedia
# ==============================================================================
try:
    from sync_master import MasterDataSync
except ImportError:
    print("PERINGATAN: 'sync_master.py' tidak ditemukan. Menggunakan mode Offline/Dummy.")
    class MasterDataSync:
        def load_data(self):
            # Return dataframe kosong sebagai fallback agar aplikasi tetap bisa dijalankan
            return pd.DataFrame(columns=['barcode', 'judul', 'harga']), False

class KasirApp:
    def __init__(self, root):
        self.root = root
        # UPDATE VERSI v1.4.1
        self.root.title("Takom Kasir v1.4.1 - Enhanced Stability & Performance")

        # AUTO LAUNCH FULL SCREEN (MAXIMIZED)
        self.root.state('zoomed')

        try:
            self.root.iconbitmap("logo.ico")
        except:
            pass

        self.sync_tool = MasterDataSync()
        self.df_master = pd.DataFrame()
        self.is_online = False
        self.is_loading = True
        
        # File & Sheet Target
        self.target_file_path = ""
        self.wb_target = None
        
        # Cache untuk performa laporan (Mencegah baca Excel berulang)
        self.cached_report_df = None
        self.cached_sheet_name = None
        
        # Konfigurasi Nota Advance
        self.config_nota = {
            "nama_toko": "TOKO KASIR TAKOM",
            "sub_nama": "Pusat Grosir & Eceran Terlengkap",
            "alamat_toko": "Jl. Raya Toko No. 88, Kota Anda",
            "telepon_toko": "Telp: 0812-3456-7890",
            "nama_kasir": "Admin", # BARU: Nama kasir dinamis
            "info_tambahan": "Terima Kasih Atas Kunjungan Anda",
            "pesan_penutup": "Barang yang sudah dibeli tidak dapat ditukar/dikembalikan.",
            "lebar_kertas": 48,
            "ukuran_font": 10,
            "tampilkan_logo_teks": True,
            "tampilkan_footer": True,
            "tampilkan_telp": True
        }
        self.load_config_nota_json()

        # Keranjang Transaksi
        self.cart = []
        self.no_trx_counter = 1

        # Bangun Styling Kustom
        self.setup_custom_styles()

        # Bangun Struktur UI Utama
        self.setup_main_layout()
        
        # Mulai load data master secara asinkron
        threading.Thread(target=self.load_master_data_async, daemon=True).start()

    def load_config_nota_json(self):
        """Memuat konfigurasi kustom nota dari file lokal jika ada"""
        try:
            if os.path.exists("config_nota.json"):
                with open("config_nota.json", "r", encoding="utf-8") as f:
                    data = json.load(f)
                    # Update hanya key yang ada, pertahankan default untuk key baru
                    for key, value in data.items():
                        self.config_nota[key] = value
        except Exception as e:
            print("Gagal memuat config nota:", e)

    def save_config_nota_json(self):
        """Menyimpan konfigurasi kustom nota ke file lokal"""
        try:
            with open("config_nota.json", "w", encoding="utf-8") as f:
                json.dump(self.config_nota, f, indent=4)
        except Exception as e:
            print("Gagal menyimpan config nota:", e)

    def setup_custom_styles(self):
        style = ttk.Style()
        style.theme_use('clam')
        style.configure(".", background="#ece9e8", fieldbackground="#ece9e8")
        
        style.configure("DarkReport.Treeview", 
                        background="#ffffff", 
                        foreground="#000000", 
                        fieldbackground="#ffffff",
                        rowheight=26,
                        font=("Arial", 10))
        
        style.configure("DarkReport.Treeview.Heading", 
                        background="#d4d0c8", 
                        foreground="#000000", 
                        font=("Arial", 10, "bold"))
        
        style.map("DarkReport.Treeview", 
                  background=[('selected', '#316ac5')], 
                  foreground=[('selected', '#ffffff')])

    def center_popup(self, popup, width, height):
        popup.update_idletasks()
        root_x = self.root.winfo_x()
        root_y = self.root.winfo_y()
        root_w = self.root.winfo_width()
        root_h = self.root.winfo_height()

        x = root_x + (root_w // 2) - (width // 2)
        y = root_y + (root_h // 2) - (height // 2)
        popup.geometry(f'{width}x{height}+{max(0, x)}+{max(0, y)}')

    def create_rounded_button(self, parent, text, command, bg_color, fg_color):
        """Membuat tombol rounded yang bisa di-update warnanya tanpa di-destroy"""
        canvas = tk.Canvas(parent, width=170, height=36, bg="#2c3e50", highlightthickness=0, cursor="hand2")
        
        def draw_button(color, text_color):
            canvas.delete("all")
            x1, y1, x2, y2, r = 2, 2, 168, 34, 8
            canvas.create_arc(x1, y1, x1 + 2*r, y1 + 2*r, start=90, extent=90, fill=color, outline=color)
            canvas.create_arc(x2 - 2*r, y1, x2, y1 + 2*r, start=0, extent=90, fill=color, outline=color)
            canvas.create_arc(x1, y2 - 2*r, x1 + 2*r, y2, start=180, extent=90, fill=color, outline=color)
            canvas.create_arc(x2 - 2*r, y2 - 2*r, x2, y2, start=270, extent=90, fill=color, outline=color)
            canvas.create_rectangle(x1 + r, y1, x2 - r, y2, fill=color, outline=color)
            canvas.create_rectangle(x1, y1 + r, x2, y2 - r, fill=color, outline=color)
            canvas.create_text(85, 18, text=text, fill=text_color, font=("Arial", 9, "bold"))

        # Simpan referensi fungsi draw agar bisa dipanggil dari luar untuk update warna
        canvas.draw = draw_button
        canvas.draw(bg_color, fg_color)
        
        canvas.bind("<Button-1>", lambda e: command())
        canvas.bind("<Enter>", lambda e: canvas.draw("#16a085" if bg_color == "#1abc9c" else "#34495e", fg_color))
        canvas.bind("<Leave>", lambda e: canvas.draw(bg_color, fg_color))
        return canvas

    def setup_main_layout(self):
        self.toolbar_frame = tk.Frame(self.root, bg="#2c3e50", height=50)
        self.toolbar_frame.pack(side=tk.TOP, fill=tk.X)
        self.toolbar_frame.pack_propagate(False)

        # Buat tombol SEKALI SAJA, simpan referensinya
        self.btn_transaksi = self.create_rounded_button(
            self.toolbar_frame, "🛒 Transaksi Kasir", 
            lambda: self.switch_tab("transaksi"), "#1abc9c", "white"
        )
        self.btn_transaksi.pack(side=tk.LEFT, padx=10, pady=7)

        self.btn_laporan = self.create_rounded_button(
            self.toolbar_frame, "📊 Laporan & Sheet", 
            lambda: self.switch_tab("laporan"), "#34495e", "white"
        )
        self.btn_laporan.pack(side=tk.LEFT, padx=5, pady=7)

        self.btn_custom_nota = ttk.Button(
            self.toolbar_frame, text="⚙️ Advanced Custom Nota", 
            command=self.buka_popup_custom_nota_advance
        )
        self.btn_custom_nota.pack(side=tk.RIGHT, padx=15, pady=8)

        self.container = ttk.Frame(self.root)
        self.container.pack(side=tk.TOP, fill=tk.BOTH, expand=True)

        self.frame_transaksi = ttk.Frame(self.container)
        self.frame_laporan = ttk.Frame(self.container)

        self.setup_halaman_transaksi(self.frame_transaksi)
        self.setup_halaman_laporan(self.frame_laporan)
        self.switch_tab("transaksi")

    def switch_tab(self, tab_name):
        """Optimasi: Hanya ubah warna tombol, jangan destroy/recreate"""
        if tab_name == "transaksi":
            self.frame_laporan.pack_forget()
            self.frame_transaksi.pack(fill=tk.BOTH, expand=True)
            
            # Update warna tombol
            self.btn_transaksi.draw("#1abc9c", "white")
            self.btn_laporan.draw("#34495e", "white")
            
            self.ent_search.focus_force()
            
        elif tab_name == "laporan":
            self.frame_transaksi.pack_forget()
            self.frame_laporan.pack(fill=tk.BOTH, expand=True)
            
            # Update warna tombol
            self.btn_transaksi.draw("#34495e", "white")
            self.btn_laporan.draw("#1abc9c", "white")
            
            self.refresh_data_laporan()

    def buka_popup_custom_nota_advance(self):
        popup = tk.Toplevel(self.root)
        popup.title("Advanced Custom Format Nota & Struk Kasir")
        self.center_popup(popup, 840, 650) # Tinggi ditambah sedikit untuk field kasir
        popup.transient(self.root)
        popup.grab_set()
        popup.bind("<Escape>", lambda e: popup.destroy())

        main_split = ttk.Frame(popup)
        main_split.pack(fill="both", expand=True, padx=10, pady=10)

        frame_form = ttk.LabelFrame(main_split, text=" Konfigurasi Lengkap Struk & Toko ")
        frame_form.pack(side="left", fill="both", expand=True, padx=5, pady=5)

        canvas_form = tk.Canvas(frame_form, bg="#ece9e8", highlightthickness=0)
        scrollbar_form = ttk.Scrollbar(frame_form, orient="vertical", command=canvas_form.yview)
        scrollable_inner = ttk.Frame(canvas_form)

        scrollable_inner.bind(
            "<Configure>",
            lambda e: canvas_form.configure(scrollregion=canvas_form.bbox("all"))
        )
        canvas_form.create_window((0, 0), window=scrollable_inner, anchor="nw")
        canvas_form.configure(yscrollcommand=scrollbar_form.set)

        scrollbar_form.pack(side="right", fill="y")
        canvas_form.pack(side="left", fill="both", expand=True)

        ttk.Label(scrollable_inner, text="Nama Toko / Header Utama:", font=("Arial", 9, "bold")).pack(anchor="w", padx=10, pady=(6, 2))
        ent_nama = ttk.Entry(scrollable_inner, width=38)
        ent_nama.insert(0, self.config_nota["nama_toko"])
        ent_nama.pack(anchor="w", padx=10)

        ttk.Label(scrollable_inner, text="Sub Nama / Slogan Toko:", font=("Arial", 9, "bold")).pack(anchor="w", padx=10, pady=(6, 2))
        ent_sub = ttk.Entry(scrollable_inner, width=38)
        ent_sub.insert(0, self.config_nota["sub_nama"])
        ent_sub.pack(anchor="w", padx=10)

        ttk.Label(scrollable_inner, text="Alamat Toko:", font=("Arial", 9, "bold")).pack(anchor="w", padx=10, pady=(6, 2))
        ent_alamat = ttk.Entry(scrollable_inner, width=38)
        ent_alamat.insert(0, self.config_nota["alamat_toko"])
        ent_alamat.pack(anchor="w", padx=10)

        ttk.Label(scrollable_inner, text="Nomor Telepon / Kontak:", font=("Arial", 9, "bold")).pack(anchor="w", padx=10, pady=(6, 2))
        ent_telp = ttk.Entry(scrollable_inner, width=38)
        ent_telp.insert(0, self.config_nota["telepon_toko"])
        ent_telp.pack(anchor="w", padx=10)

        # BARU: Field Nama Kasir Default
        ttk.Label(scrollable_inner, text="Nama Kasir Default:", font=("Arial", 9, "bold")).pack(anchor="w", padx=10, pady=(6, 2))
        ent_kasir = ttk.Entry(scrollable_inner, width=38)
        ent_kasir.insert(0, self.config_nota.get("nama_kasir", "Admin"))
        ent_kasir.pack(anchor="w", padx=10)

        ttk.Label(scrollable_inner, text="Pesan Footer / Terima Kasih:", font=("Arial", 9, "bold")).pack(anchor="w", padx=10, pady=(6, 2))
        ent_footer = ttk.Entry(scrollable_inner, width=38)
        ent_footer.insert(0, self.config_nota["info_tambahan"])
        ent_footer.pack(anchor="w", padx=10)

        ttk.Label(scrollable_inner, text="Catatan Syarat & Ketentuan Bawah:", font=("Arial", 9, "bold")).pack(anchor="w", padx=10, pady=(6, 2))
        ent_penutup = ttk.Entry(scrollable_inner, width=38)
        ent_penutup.insert(0, self.config_nota["pesan_penutup"])
        ent_penutup.pack(anchor="w", padx=10)

        row_num = ttk.Frame(scrollable_inner)
        row_num.pack(anchor="w", padx=10, pady=10)
        
        ttk.Label(row_num, text="Ukuran Font:").pack(side="left", padx=(0, 2))
        spin_font = ttk.Spinbox(row_num, from_=8, to=14, width=4)
        spin_font.set(self.config_nota["ukuran_font"])
        spin_font.pack(side="left", padx=2)

        ttk.Label(row_num, text="Lebar Kertas:").pack(side="left", padx=(10, 2))
        spin_lebar = ttk.Spinbox(row_num, from_=32, to=80, width=4)
        spin_lebar.set(self.config_nota["lebar_kertas"])
        spin_lebar.pack(side="left", padx=2)

        logo_var = tk.BooleanVar(value=self.config_nota["tampilkan_logo_teks"])
        chk_logo = ttk.Checkbutton(scrollable_inner, text="Tampilkan Header Nama Toko", variable=logo_var)
        chk_logo.pack(anchor="w", padx=10, pady=2)

        telp_var = tk.BooleanVar(value=self.config_nota["tampilkan_telp"])
        chk_telp = ttk.Checkbutton(scrollable_inner, text="Tampilkan Info Telepon", variable=telp_var)
        chk_telp.pack(anchor="w", padx=10, pady=2)

        footer_var = tk.BooleanVar(value=self.config_nota["tampilkan_footer"])
        chk_footer = ttk.Checkbutton(scrollable_inner, text="Tampilkan Catatan Kaki / Footer", variable=footer_var)
        chk_footer.pack(anchor="w", padx=10, pady=6)

        frame_preview = ttk.LabelFrame(main_split, text=" Live Preview Tampilan Nota Advance ")
        frame_preview.pack(side="right", fill="both", expand=True, padx=5, pady=5)

        txt_preview = tk.Text(frame_preview, width=40, height=27, font=("Courier New", 9), bg="#fcfcfc")
        txt_preview.pack(side="top", fill="both", expand=True, padx=5, pady=5)

        def update_preview(event=None):
            try:
                lebar = int(spin_lebar.get())
                nama = ent_nama.get()
                sub = ent_sub.get()
                alamat = ent_alamat.get()
                telp = ent_telp.get()
                kasir = ent_kasir.get()
                footer = ent_footer.get()
                penutup = ent_penutup.get()
                show_logo = logo_var.get()
                show_telp = telp_var.get()
                show_f = footer_var.get()
                f_size = spin_font.get()

                line = "-" * min(lebar, 48)
                preview_lines = []
                if show_logo:
                    preview_lines.append(nama.center(min(lebar, 48)))
                    preview_lines.append(sub.center(min(lebar, 48)))
                
                preview_lines.append(alamat.center(min(lebar, 48)))
                if show_telp:
                    preview_lines.append(telp.center(min(lebar, 48)))
                
                preview_lines.append(line)
                preview_lines.append(f"No. Trx  : #1025           Kasir : {kasir}".ljust(min(lebar, 48)))
                preview_lines.append("Waktu   : 12/09/2026 16:22".ljust(min(lebar, 48)))
                preview_lines.append(line)
                preview_lines.append("Rosario Kayu Wangi Eksklusif".ljust(min(lebar, 48)))
                preview_lines.append("  2 x @25,000                Rp. 50,000".ljust(min(lebar, 48)))
                preview_lines.append("Buku Doa Katolik Saku".ljust(min(lebar, 48)))
                preview_lines.append("  1 x @15,000                Rp. 15,000".ljust(min(lebar, 48)))
                preview_lines.append(line)
                preview_lines.append("TOTAL BELANJA              Rp. 65,000".ljust(min(lebar, 48)))
                preview_lines.append("PEMBAYARAN (TUNAI)         Rp. 70,000".ljust(min(lebar, 48)))
                preview_lines.append("KEMBALIAN                   Rp. 5,000".ljust(min(lebar, 48)))
                preview_lines.append(line)
                
                if show_f:
                    preview_lines.append(footer.center(min(lebar, 48)))
                    preview_lines.append(penutup.center(min(lebar, 48)))
                    
                preview_lines.append(f"[Font: {f_size}pt | Lebar: {lebar}kolom]".center(min(lebar, 48)))

                txt_preview.config(state="normal")
                txt_preview.delete("1.0", tk.END)
                txt_preview.insert("1.0", "\n".join(preview_lines))
                txt_preview.config(state="disabled")
            except:
                pass

        for widget_entry in [ent_nama, ent_sub, ent_alamat, ent_telp, ent_kasir, ent_footer, ent_penutup, spin_font, spin_lebar]:
            widget_entry.bind("<KeyRelease>", update_preview)
        
        chk_logo.configure(command=update_preview)
        chk_telp.configure(command=update_preview)
        chk_footer.configure(command=update_preview)
        update_preview()

        def simpan_pengaturan():
            try:
                self.config_nota["nama_toko"] = ent_nama.get().strip()
                self.config_nota["sub_nama"] = ent_sub.get().strip()
                self.config_nota["alamat_toko"] = ent_alamat.get().strip()
                self.config_nota["telepon_toko"] = ent_telp.get().strip()
                self.config_nota["nama_kasir"] = ent_kasir.get().strip() or "Admin"
                self.config_nota["info_tambahan"] = ent_footer.get().strip()
                self.config_nota["pesan_penutup"] = ent_penutup.get().strip()
                self.config_nota["lebar_kertas"] = int(spin_lebar.get())
                self.config_nota["ukuran_font"] = int(spin_font.get())
                self.config_nota["tampilkan_logo_teks"] = logo_var.get()
                self.config_nota["tampilkan_telp"] = telp_var.get()
                self.config_nota["tampilkan_footer"] = footer_var.get()

                self.save_config_nota_json()
                messagebox.showinfo("Sukses", "Format advanced kustom nota berhasil disimpan!", parent=popup)
                popup.destroy()
            except Exception as e:
                messagebox.showerror("Error", f"Gagal menyimpan format:\n{str(e)}", parent=popup)

        ttk.Button(popup, text="SIMPAN FORMAT ADVANCE", command=simpan_pengaturan).pack(side="bottom", pady=12)

    def setup_halaman_transaksi(self, parent):
        top_container = ttk.Frame(parent)
        top_container.pack(fill="x", padx=10, pady=5)

        frame_config = ttk.LabelFrame(top_container, text=" Pengaturan File & Sheet Target ")
        frame_config.pack(side="left", fill="both", expand=True)

        row_master = ttk.Frame(frame_config)
        row_master.pack(fill="x", padx=5, pady=6)
        
        self.lbl_status = ttk.Label(row_master, text="[ 🔄 Connecting Database... Mohon Tunggu ]", font=("Arial", 9, "bold"), foreground="blue")
        self.lbl_status.pack(side="left", padx=5)

        row_target = ttk.Frame(frame_config)
        row_target.pack(fill="x", padx=5, pady=6)

        btn_target = ttk.Button(row_target, text="Pilih File Target Output", command=self.pilih_target_file)
        btn_target.pack(side="left", padx=5)

        self.lbl_target_path = ttk.Label(row_target, text="Belum dipilih", foreground="red", font=("Arial", 9, "bold"))
        self.lbl_target_path.pack(side="left", padx=5)

        ttk.Label(row_target, text="Pilih Sheet Target:").pack(side="left", padx=(20, 2))
        self.combo_sheet = ttk.Combobox(row_target, width=15, state="readonly")
        self.combo_sheet.pack(side="left", padx=2)
        self.combo_sheet.bind("<<ComboboxSelected>>", self.on_sheet_changed)

        btn_add_sheet = ttk.Button(row_target, text="+ Sheet", width=8, command=self.tambah_sheet_baru)
        btn_add_sheet.pack(side="left", padx=5)

        info_struktur = ttk.Label(
            frame_config, 
            text="📌 Struktur Kolom Excel Baku (A s.d. K): [A] No | [B] Kode Produk | [C] Judul | [D] Jumlah | [E] Harga Satuan | [F] Total Per Produk | [G-H] Metode (Tunai/Non-Tunai) | [I] Diskon | [J] Member | [K] Waktu", 
            font=("Arial", 8, "italic"),
            foreground="darkslategray"
        )
        info_struktur.pack(anchor="w", padx=10, pady=(4, 8))

        scan_frame = ttk.LabelFrame(parent, text=" Area Scan / Cari Produk ")
        scan_frame.pack(fill="x", padx=10, pady=5)

        ttk.Label(scan_frame, text="Cari Barcode / Judul:", font=("Arial", 10, "bold")).pack(side="left", padx=5)
        self.ent_search = ttk.Entry(scan_frame, font=("Arial", 11))
        self.ent_search.pack(side="left", fill="x", expand=True, padx=5, pady=5)
        self.ent_search.bind("<Return>", self.proses_scan_trigger)
        self.ent_search.bind("<Down>", self.fokus_ke_tabel)
        self.ent_search.focus()

        table_frame = ttk.Frame(parent)
        table_frame.pack(fill="both", expand=True, padx=10, pady=5)

        columns = ("no", "barcode", "judul", "jumlah", "harga_normal", "diskon", "harga_akhir", "total")
        self.tree = ttk.Treeview(table_frame, columns=columns, show="headings", selectmode="browse")

        self.tree.heading("no", text="No")
        self.tree.heading("barcode", text="Barcode")
        self.tree.heading("judul", text="Judul")
        self.tree.heading("jumlah", text="Jumlah")
        self.tree.heading("harga_normal", text="Harga Normal")
        self.tree.heading("diskon", text="Diskon")
        self.tree.heading("harga_akhir", text="Harga Akhir")
        self.tree.heading("total", text="Total")

        self.tree.column("no", width=40, anchor="center")
        self.tree.column("barcode", width=120)
        self.tree.column("judul", width=380)
        self.tree.column("jumlah", width=60, anchor="center")
        self.tree.column("harga_normal", width=100, anchor="e")
        self.tree.column("diskon", width=80, anchor="center")
        self.tree.column("harga_akhir", width=100, anchor="e")
        self.tree.column("total", width=120, anchor="e")

        self.tree.pack(fill="both", expand=True)

        self.tree.bind("<Double-1>", self.buka_popup_edit_qty_langsung)
        self.tree.bind("<Delete>", self.hapus_item_terpilih)
        self.tree.bind("<BackSpace>", self.hapus_item_terpilih)
        self.tree.bind("<Escape>", lambda event: self.ent_search.focus_force())
        self.root.bind("<F5>", lambda event: self.buka_popup_pembayaran())

        bottom_frame = ttk.Frame(parent)
        bottom_frame.pack(fill="x", padx=10, pady=10)

        info_summary_frame = ttk.Frame(bottom_frame)
        info_summary_frame.pack(side="top", fill="x", pady=5)

        self.lbl_total_qty = ttk.Label(info_summary_frame, text="TOTAL ITEM / QTY: 0 Pcs", font=("Arial", 12, "bold"), foreground="darkgreen")
        self.lbl_total_qty.pack(side="left", padx=10)

        self.lbl_total = ttk.Label(info_summary_frame, text="TOTAL BELANJA: Rp. 0", font=("Arial", 16, "bold"), foreground="blue")
        self.lbl_total.pack(side="right", padx=10)

        btn_row = ttk.Frame(bottom_frame)
        btn_row.pack(side="bottom", fill="x", pady=5)

        btn_delete_item = ttk.Button(btn_row, text="Hapus Item Selected (Del)", command=self.hapus_item_terpilih)
        btn_delete_item.pack(side="left", padx=5)

        btn_reset = ttk.Button(btn_row, text="Reset Transaksi Saat Ini", command=self.reset_transaksi)
        btn_reset.pack(side="right", padx=5)

        btn_finish = ttk.Button(btn_row, text="SELESAI & SIMPAN TRANSAKSI (F5)", command=self.buka_popup_pembayaran)
        btn_finish.pack(side="right", padx=5)

    def setup_halaman_laporan(self, parent):
        frame_top_rep = ttk.Frame(parent)
        frame_top_rep.pack(fill="x", padx=10, pady=10)

        lbl_rep_title = ttk.Label(frame_top_rep, text="Pilih Sheet Laporan:", font=("Arial", 10, "bold"))
        lbl_rep_title.pack(side="left", padx=5)
        
        self.combo_sheet_report = ttk.Combobox(frame_top_rep, width=20, state="readonly")
        self.combo_sheet_report.pack(side="left", padx=5)
        self.combo_sheet_report.bind("<<ComboboxSelected>>", self.muat_tabel_laporan_excel)

        btn_refresh = ttk.Button(frame_top_rep, text="🔄 Refresh Data", command=self.force_refresh_laporan)
        btn_refresh.pack(side="left", padx=5)

        btn_setoran = ttk.Button(frame_top_rep, text="📊 Setoran Harian", command=self.proses_setoran_harian)
        btn_setoran.pack(side="left", padx=15)

        btn_reprint = ttk.Button(frame_top_rep, text="🖨️ Re-Print Nota Terpilih", command=self.reprint_nota_dari_sheet)
        btn_reprint.pack(side="left", padx=5)

        self.frame_summary_box = tk.Frame(parent, bg="#dcd6d0", bd=2, relief="groove")
        self.frame_summary_box.pack(fill="x", padx=10, pady=5)

        self.lbl_setoran_tunai = tk.Label(self.frame_summary_box, text="Total Tunai: Rp. 0", font=("Arial", 11, "bold"), bg="#dcd6d0", fg="#111111")
        self.lbl_setoran_tunai.pack(side="left", padx=15, pady=8)

        self.lbl_setoran_nontunai = tk.Label(self.frame_summary_box, text="Total Non-Tunai: Rp. 0", font=("Arial", 11, "bold"), bg="#dcd6d0", fg="#111111")
        self.lbl_setoran_nontunai.pack(side="left", padx=15, pady=8)

        self.lbl_setoran_grand = tk.Label(self.frame_summary_box, text="GRAND TOTAL SETORAN: Rp. 0", font=("Arial", 12, "bold"), bg="#dcd6d0", fg="#0000aa")
        self.lbl_setoran_grand.pack(side="right", padx=15, pady=8)

        table_frame_rep = ttk.Frame(parent)
        table_frame_rep.pack(fill="both", expand=True, padx=10, pady=5)

        self.report_tree = ttk.Treeview(table_frame_rep, show="headings", style="DarkReport.Treeview", selectmode="browse")
        
        sb_y = ttk.Scrollbar(table_frame_rep, orient="vertical", command=self.report_tree.yview)
        sb_x = ttk.Scrollbar(table_frame_rep, orient="horizontal", command=self.report_tree.xview)
        self.report_tree.configure(yscrollcommand=sb_y.set, xscrollcommand=sb_x.set)

        sb_y.pack(side="right", fill="y")
        sb_x.pack(side="bottom", fill="x")
        self.report_tree.pack(side="left", fill="both", expand=True)

    def force_refresh_laporan(self):
        """Memaksa pembacaan ulang file Excel (menghapus cache)"""
        self.cached_report_df = None
        self.cached_sheet_name = None
        self.refresh_data_laporan()
        messagebox.showinfo("Info", "Data laporan berhasil diperbarui dari file Excel.")

    def refresh_data_laporan(self):
        if not self.target_file_path or not self.wb_target:
            self.combo_sheet_report['values'] = []
            return
        
        sheets = self.wb_target.sheetnames
        self.combo_sheet_report['values'] = sheets
        current_active = self.combo_sheet.get()
        if current_active in sheets:
            self.combo_sheet_report.set(current_active)
        elif sheets:
            self.combo_sheet_report.current(0)
            
        self.muat_tabel_laporan_excel()

    def muat_tabel_laporan_excel(self, event=None):
        for item in self.report_tree.get_children():
            self.report_tree.delete(item)

        if not self.target_file_path or not self.wb_target:
            return

        sheet_name = self.combo_sheet_report.get()
        if not sheet_name or sheet_name not in self.wb_target.sheetnames:
            return

        # OPTIMASI: Gunakan cache jika sheet yang sama belum berubah
        if self.cached_report_df is not None and self.cached_sheet_name == sheet_name:
            df_raw = self.cached_report_df
        else:
            try:
                df_raw = pd.read_excel(self.target_file_path, sheet_name=sheet_name, header=None)
                self.cached_report_df = df_raw
                self.cached_sheet_name = sheet_name
            except Exception as e:
                messagebox.showerror("Error", f"Gagal membaca file Excel:\n{str(e)}")
                return

        if df_raw.empty or len(df_raw) < 1:
            return

        row1 = df_raw.iloc[0].values if len(df_raw) > 0 else []
        row2 = df_raw.iloc[1].values if len(df_raw) > 1 else []

        headers = []
        for i in range(df_raw.shape[1]):
            val1 = str(row1[i]).strip() if i < len(row1) and pd.notna(row1[i]) else ""
            val2 = str(row2[i]).strip() if i < len(row2) and pd.notna(row2[i]) else ""
            
            if val1 and val2 and val1 != val2:
                header_text = f"{val1} - {val2}"
            elif val1:
                header_text = val1
            elif val2:
                header_text = val2
            else:
                header_text = f"Kolom {i+1}"
            
            if header_text.lower() in ["no transaksi", "no. transaksi"]:
                header_text = "No"

            headers.append(header_text)

        cols = [f"col_{i}" for i in range(df_raw.shape[1])]
        self.report_tree["columns"] = cols
        
        for idx, col in enumerate(cols):
            h_text = headers[idx] if idx < len(headers) else col
            self.report_tree.heading(col, text=h_text)
            
            max_len = len(h_text)
            for r_idx in range(1, min(50, len(df_raw))): # Batasi sampling untuk performa
                cell_val = str(df_raw.iloc[r_idx, idx]) if pd.notna(df_raw.iloc[r_idx, idx]) else ""
                if len(cell_val) > max_len:
                    max_len = len(cell_val)
            
            col_width = 45 if idx == 0 else max(max_len * 9, 90)
            self.report_tree.column(col, width=col_width, anchor="w", stretch=False)

        for r_idx in range(2, len(df_raw)):
            row_vals = df_raw.iloc[r_idx].values
            vals = [str(val) if pd.notna(val) else "" for val in row_vals]
            while len(vals) < len(cols):
                vals.append("")
            self.report_tree.insert("", "end", values=vals[:len(cols)])

        self.hitung_dan_tampilkan_setoran(df_raw)

    def hitung_dan_tampilkan_setoran(self, df_raw):
        total_tunai = 0.0
        total_nontunai = 0.0

        for r_idx in range(2, len(df_raw)):
            try:
                val_a = str(df_raw.iloc[r_idx, 0]).strip().upper()
                if "SETORAN" in val_a or "TOTAL" in val_a or "TUTUP" in val_a:
                    continue
                
                val_g = df_raw.iloc[r_idx, 6] if df_raw.shape[1] > 6 else 0
                if pd.notna(val_g):
                    total_tunai += float(str(val_g).replace(",", ""))
            except:
                pass

            try:
                val_a = str(df_raw.iloc[r_idx, 0]).strip().upper()
                if "SETORAN" in val_a or "TOTAL" in val_a or "TUTUP" in val_a:
                    continue

                val_h = df_raw.iloc[r_idx, 7] if df_raw.shape[1] > 7 else 0
                if pd.notna(val_h):
                    total_nontunai += float(str(val_h).replace(",", ""))
            except:
                pass

        grand_setoran = total_tunai + total_nontunai

        self.lbl_setoran_tunai.config(text=f"Total Tunai: Rp. {total_tunai:,.0f}")
        self.lbl_setoran_nontunai.config(text=f"Total Non-Tunai: Rp. {total_nontunai:,.0f}")
        self.lbl_setoran_grand.config(text=f"GRAND TOTAL SETORAN: Rp. {grand_setoran:,.0f}")

    def reprint_nota_dari_sheet(self):
        selected_items = self.report_tree.selection()
        if not selected_items:
            messagebox.showwarning("Peringatan", "Pilih baris transaksi pada tabel laporan terlebih dahulu untuk di-reprint!")
            return

        sheet_name = self.combo_sheet_report.get()
        if not sheet_name or not self.target_file_path:
            return

        try:
            # Gunakan cache jika ada, jika tidak baca ulang
            if self.cached_report_df is not None and self.cached_sheet_name == sheet_name:
                df_raw = self.cached_report_df
            else:
                df_raw = pd.read_excel(self.target_file_path, sheet_name=sheet_name, header=None)

            selected_row_values = self.report_tree.item(selected_items[0], "values")
            if not selected_row_values:
                return

            target_no_trx = None
            clicked_excel_row_idx = -1

            # Pencarian yang lebih robust
            for r_idx in range(2, len(df_raw)):
                try:
                    r_vals = [str(df_raw.iloc[r_idx, c]) if pd.notna(df_raw.iloc[r_idx, c]) else "" for c in range(min(3, df_raw.shape[1]))]
                    if r_vals and str(r_vals[0]).strip() == str(selected_row_values[0]).strip() and str(r_vals[1]).strip() == str(selected_row_values[1]).strip():
                        clicked_excel_row_idx = r_idx
                        break
                except:
                    continue
            
            if clicked_excel_row_idx == -1:
                for i, child_id in enumerate(self.report_tree.get_children()):
                    if child_id == selected_items[0]:
                        clicked_excel_row_idx = i + 2
                        break

            if clicked_excel_row_idx == -1:
                messagebox.showerror("Error", "Gagal mendeteksi baris transaksi pada Excel.")
                return

            current_r = clicked_excel_row_idx
            while current_r >= 2:
                try:
                    val_a = str(df_raw.iloc[current_r, 0]).strip()
                    if val_a and val_a.replace(".", "").isdigit():
                        target_no_trx = val_a
                        break
                except:
                    pass
                current_r -= 1

            if not target_no_trx:
                target_no_trx = "1"

            items_reprint = []
            metode_reprint = "TUNAI"
            waktu_reprint = datetime.now().strftime("%H:%M:%S")
            member_reprint = ""

            start_scan = current_r
            while start_scan < len(df_raw):
                try:
                    val_a_check = str(df_raw.iloc[start_scan, 0]).strip()
                    if val_a_check and val_a_check.replace(".", "").isdigit() and val_a_check != target_no_trx and start_scan != current_r:
                        break
                    
                    barcode = str(df_raw.iloc[start_scan, 1]) if pd.notna(df_raw.iloc[start_scan, 1]) else ""
                    judul = str(df_raw.iloc[start_scan, 2]) if pd.notna(df_raw.iloc[start_scan, 2]) else ""
                    qty = float(df_raw.iloc[start_scan, 3]) if pd.notna(df_raw.iloc[start_scan, 3]) else 1
                    harga = float(df_raw.iloc[start_scan, 4]) if pd.notna(df_raw.iloc[start_scan, 4]) else 0
                    subtotal = float(df_raw.iloc[start_scan, 5]) if pd.notna(df_raw.iloc[start_scan, 5]) else (qty * harga)
                    
                    val_g = df_raw.iloc[start_scan, 6] if df_raw.shape[1] > 6 else None
                    val_h = df_raw.iloc[start_scan, 7] if df_raw.shape[1] > 7 else None
                    if pd.notna(val_g) and str(val_g).strip() != "":
                        metode_reprint = "TUNAI"
                    elif pd.notna(val_h) and str(val_h).strip() != "":
                        metode_reprint = "NON TUNAI"

                    diskon = float(df_raw.iloc[start_scan, 8]) if df_raw.shape[1] > 8 and pd.notna(df_raw.iloc[start_scan, 8]) else 0.0
                    member_val = str(df_raw.iloc[start_scan, 9]) if df_raw.shape[1] > 9 and pd.notna(df_raw.iloc[start_scan, 9]) else ""
                    if member_val:
                        member_reprint = member_val

                    time_val = str(df_raw.iloc[start_scan, 10]) if df_raw.shape[1] > 10 and pd.notna(df_raw.iloc[start_scan, 10]) else ""
                    if time_val:
                        waktu_reprint = time_val

                    if barcode or judul:
                        items_reprint.append({
                            "judul": judul, "jumlah": int(qty), "harga_akhir": harga, "total": subtotal, "diskon": diskon
                        })
                except:
                    pass

                start_scan += 1

            if not items_reprint:
                messagebox.showwarning("Peringatan", "Tidak ada item valid yang ditemukan untuk transaksi ini.")
                return

            grand_total_rep = sum(i['total'] for i in items_reprint)
            self.cetak_nota_thermal_80mm_custom(
                no_trx=target_no_trx,
                metode=metode_reprint,
                bayar=grand_total_rep,
                kembali=grand_total_rep,
                catatan_member=member_reprint,
                waktu_str=waktu_reprint,
                items_source=items_reprint,
                is_reprint=True
            )
            messagebox.showinfo("Sukses Re-Print", f"Nota Transaksi No. #{target_no_trx} berhasil dicetak ulang (Re-Print)!")

        except Exception as e:
            messagebox.showerror("Error Re-Print", f"Gagal memproses re-print nota:\n{str(e)}")

    def proses_setoran_harian(self):
        if not self.target_file_path or not self.wb_target:
            messagebox.showwarning("Peringatan", "Pilih File Target Output terlebih dahulu!")
            return

        sheet_name = self.combo_sheet_report.get()
        if not sheet_name or sheet_name not in self.wb_target.sheetnames:
            messagebox.showwarning("Peringatan", "Pilih sheet laporan yang valid!")
            return

        try:
            ws = self.wb_target[sheet_name]
            
            # Gunakan cache untuk perhitungan agar cepat
            if self.cached_report_df is not None and self.cached_sheet_name == sheet_name:
                df_raw = self.cached_report_df
            else:
                df_raw = pd.read_excel(self.target_file_path, sheet_name=sheet_name, header=None)
            
            total_tunai = 0.0
            total_nontunai = 0.0
            for r_idx in range(2, len(df_raw)):
                try:
                    val_a = str(df_raw.iloc[r_idx, 0]).strip().upper()
                    if "SETORAN" in val_a or "TOTAL" in val_a or "TUTUP" in val_a:
                        continue
                    vg = df_raw.iloc[r_idx, 6]
                    if pd.notna(vg): total_tunai += float(str(vg).replace(",", ""))
                except: pass
                try:
                    val_a = str(df_raw.iloc[r_idx, 0]).strip().upper()
                    if "SETORAN" in val_a or "TOTAL" in val_a or "TUTUP" in val_a:
                        continue
                    vh = df_raw.iloc[r_idx, 7]
                    if pd.notna(vh): total_nontunai += float(str(vh).replace(",", ""))
                except: pass
            grand_setoran = total_tunai + total_nontunai

            max_r = ws.max_row + 2
            
            ws.merge_cells(start_row=max_r, start_column=1, end_row=max_r, end_column=4)
            ws[f"A{max_r}"] = "SETORAN HARIAN (TUTUP BUKU)"
            ws[f"G{max_r}"] = total_tunai
            ws[f"H{max_r}"] = total_nontunai
            ws[f"F{max_r}"] = grand_setoran

            bold_font = Font(name="Calibri", size=11, bold=True)
            currency_format = '"Rp" #,##0'
            thin_border = Border(
                left=Side(style='thin', color='000000'), right=Side(style='thin', color='000000'),
                top=Side(style='thin', color='000000'), bottom=Side(style='thin', color='000000')
            )

            for col_l in ["A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K"]:
                cell = ws[f"{col_l}{max_r}"]
                cell.font = bold_font
                cell.border = thin_border
                if col_l in ["F", "G", "H"]:
                    cell.number_format = currency_format

            self.wb_target.save(self.target_file_path)
            
            # Invalidasi cache karena file Excel telah dimodifikasi
            self.cached_report_df = None
            self.cached_sheet_name = None
            
            self.muat_tabel_laporan_excel()
            messagebox.showinfo("Sukses", f"Setoran Harian berhasil ditutup & disimpan ke Excel!")
        except Exception as e:
            messagebox.showerror("Error", f"Gagal memproses setoran harian:\n{str(e)}")

    def buka_popup_edit_qty_langsung(self, event=None):
        selected_items = self.tree.selection()
        if not selected_items:
            return
        
        item_id = selected_items[0]
        item_index = self.tree.index(item_id)
        if not (0 <= item_index < len(self.cart)):
            return

        current_cart_item = self.cart[item_index]
        popup = tk.Toplevel(self.root)
        popup.title("Ubah Jumlah Qty")
        self.center_popup(popup, 320, 160)
        popup.transient(self.root)
        popup.grab_set()
        popup.bind("<Escape>", lambda e: popup.destroy())

        ttk.Label(popup, text=f"Produk: {current_cart_item['judul']}", font=("Arial", 9, "bold"), wraplength=300).pack(pady=(10, 5))
        
        row_q = ttk.Frame(popup)
        row_q.pack(pady=5)
        ttk.Label(row_q, text="Jumlah Baru (Qty):").pack(side="left", padx=5)
        ent_new_qty = ttk.Entry(row_q, width=10, font=("Arial", 11, "bold"))
        ent_new_qty.insert(0, str(current_cart_item['jumlah']))
        ent_new_qty.pack(side="left", padx=5)
        ent_new_qty.focus_force()
        ent_new_qty.selection_range(0, tk.END)

        def save_new_qty(event=None):
            try:
                new_q = int(ent_new_qty.get().strip())
                if new_q <= 0: raise ValueError()
                current_cart_item['jumlah'] = new_q
                current_cart_item['total'] = new_q * current_cart_item['harga_akhir']
                self.update_tabel_keranjang()
                popup.destroy()
                self.ent_search.focus_force()
            except:
                messagebox.showwarning("Peringatan", "Masukkan angka bulat positif untuk jumlah!", parent=popup)

        popup.bind("<Return>", save_new_qty)
        ttk.Button(popup, text="Simpan (ENTER)", command=save_new_qty).pack(pady=10)

    def cek_dan_buat_header_excel(self, ws):
        val_a1 = ws["A1"].value
        if str(val_a1).strip().lower() not in ["no transaksi", "no"]:
            thin_border = Border(
                left=Side(style='thin', color='000000'), right=Side(style='thin', color='000000'),
                top=Side(style='thin', color='000000'), bottom=Side(style='thin', color='000000')
            )
            header_font = Font(name="Calibri", size=11, bold=True)
            align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

            fill_standard = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
            fill_nontunai = PatternFill(start_color="BFBFBF", end_color="BFBFBF", fill_type="solid")

            single_cols = [
                ("A", "No"), ("B", "Kode Produk"), ("C", "Judul"), ("D", "Jumlah"),
                ("E", "Harga satuan"), ("F", "Total Per Produk"), ("I", "Diskon"), ("J", "Member"), ("K", "Waktu")
            ]

            for col, text in single_cols:
                ws.merge_cells(f"{col}1:{col}2")
                cell1 = ws[f"{col}1"]
                cell2 = ws[f"{col}2"]
                cell1.value = text
                cell1.font = header_font
                cell1.alignment = align_center
                cell1.fill = fill_standard
                cell2.fill = fill_standard
                cell1.border = thin_border
                cell2.border = thin_border

            ws.merge_cells("G1:H1")
            cell_g1 = ws["G1"]
            cell_g1.value = "Metode"
            cell_g1.font = header_font
            cell_g1.alignment = align_center
            cell_g1.fill = fill_standard
            ws["G1"].border = thin_border
            ws["H1"].border = thin_border

            ws["G2"].value = "TUNAI"
            ws["G2"].font = header_font
            ws["G2"].alignment = align_center
            ws["G2"].fill = fill_standard
            ws["G2"].border = thin_border

            ws["H2"].value = "NON TUNAI"
            ws["H2"].font = header_font
            ws["H2"].alignment = align_center
            ws["H2"].fill = fill_nontunai
            ws["H2"].border = thin_border

            self.wb_target.save(self.target_file_path)

    def hitung_nomor_transaksi_berikutnya(self):
        if not self.target_file_path or not self.wb_target:
            self.no_trx_counter = 1
            return
        try:
            sheet_name = self.combo_sheet.get()
            if not sheet_name or sheet_name not in self.wb_target.sheetnames:
                self.no_trx_counter = 1
                return
            ws = self.wb_target[sheet_name]
            self.cek_dan_buat_header_excel(ws)
            max_r = ws.max_row
            last_trx = 0
            for r in range(max_r, 2, -1):
                val = ws[f"A{r}"].value
                if val is not None:
                    try:
                        last_trx = int(float(str(val).strip()))
                        break
                    except ValueError:
                        continue
            self.no_trx_counter = last_trx + 1 if last_trx > 0 else 1
        except Exception as e:
            print("Gagal membaca nomor transaksi terakhir:", e)
            self.no_trx_counter = 1

    def on_sheet_changed(self, event=None):
        self.hitung_nomor_transaksi_berikutnya()

    def fokus_ke_tabel(self, event=None):
        children = self.tree.get_children()
        if children:
            self.tree.focus_set()
            if not self.tree.selection():
                self.tree.selection_set(children[0])
                self.tree.focus(children[0])

    def hapus_item_terpilih(self, event=None):
        selected_items = self.tree.selection()
        if not selected_items:
            messagebox.showwarning("Peringatan", "Pilih item di tabel yang ingin dihapus terlebih dahulu!")
            return
        selected_item = selected_items[0]
        item_index = self.tree.index(selected_item)
        if 0 <= item_index < len(self.cart):
            del self.cart[item_index]
            self.update_tabel_keranjang()
            children = self.tree.get_children()
            if children:
                new_index = min(item_index, len(children) - 1)
                self.tree.selection_set(children[new_index])
                self.tree.focus(children[new_index])
                self.tree.focus_set()
            else:
                self.ent_search.focus_force()

    def load_master_data_async(self):
        df, is_online = self.sync_tool.load_data()
        self.df_master = df
        self.is_online = is_online
        self.is_loading = False

        def update_label():
            if self.is_online:
                self.lbl_status.config(text=f"[ 🟢 ONLINE Cloud: {len(self.df_master):,} Produk ]", foreground="green")
            else:
                self.lbl_status.config(text=f"[ 🟡 OFFLINE Cache: {len(self.df_master):,} Produk ]", foreground="orange")

        self.root.after(0, update_label)

    def pilih_target_file(self):
        file_path = filedialog.askopenfilename(filetypes=[("Excel Files", "*.xlsx")])
        if file_path:
            self.target_file_path = file_path
            self.lbl_target_path.config(text=os.path.basename(file_path), foreground="green")
            try:
                self.wb_target = openpyxl.load_workbook(file_path)
                sheets = self.wb_target.sheetnames
                self.combo_sheet['values'] = sheets
                if sheets:
                    self.combo_sheet.current(0)
                
                # Reset cache karena file baru dibuka
                self.cached_report_df = None
                self.cached_sheet_name = None
                
                self.hitung_nomor_transaksi_berikutnya()
            except Exception as e:
                messagebox.showerror("Error Excel", f"Gagal membaca file target:\n{str(e)}")

    def tambah_sheet_baru(self):
        if not self.target_file_path or not self.wb_target:
            messagebox.showwarning("Peringatan", "Pilih File Target Output terlebih dahulu!")
            return
        popup = tk.Toplevel(self.root)
        popup.title("Tambah Sheet Baru")
        self.center_popup(popup, 300, 130)
        popup.transient(self.root)
        popup.grab_set()
        popup.bind("<Escape>", lambda e: popup.destroy())

        ttk.Label(popup, text="Nama Sheet Baru:").pack(pady=10)
        ent_sheet_name = ttk.Entry(popup, width=25)
        ent_sheet_name.pack(pady=5)
        ent_sheet_name.focus_force()

        def submit_sheet():
            name = ent_sheet_name.get().strip()
            if name:
                if name in self.wb_target.sheetnames:
                    messagebox.showwarning("Peringatan", "Nama sheet sudah ada!", parent=popup)
                else:
                    new_ws = self.wb_target.create_sheet(title=name)
                    self.cek_dan_buat_header_excel(new_ws)
                    self.wb_target.save(self.target_file_path)
                    self.combo_sheet['values'] = self.wb_target.sheetnames
                    self.combo_sheet.set(name)
                    
                    # Reset cache
                    self.cached_report_df = None
                    self.cached_sheet_name = None
                    
                    self.hitung_nomor_transaksi_berikutnya()
                    popup.destroy()

        popup.bind("<Return>", lambda e: submit_sheet())
        ttk.Button(popup, text="Tambah", command=submit_sheet).pack(pady=5)

    def proses_scan_trigger(self, event=None):
        if self.is_loading:
            messagebox.showinfo("Proses Loading", "Database master sedang dimuat dari cloud. Harap tunggu sebentar!")
            return
        raw_input = self.ent_search.get().strip()
        if not raw_input:
            return
        if not self.target_file_path or not self.combo_sheet.get():
            messagebox.showwarning("File Output Belum Dipilih", "Harap pilih 'File Target Output' & 'Sheet Target' terlebih dahulu!")
            return

        match_barcode = self.df_master[self.df_master['barcode'].astype(str) == raw_input]
        if not match_barcode.empty:
            item = match_barcode.iloc[0]
            self.buka_popup_input_produk(item)
            self.ent_search.delete(0, tk.END)
            return

        match_manual = self.df_master[
            self.df_master['barcode'].astype(str).str.contains(raw_input, case=False, na=False) |
            self.df_master['judul'].astype(str).str.contains(raw_input, case=False, na=False)
        ]
        if not match_manual.empty:
            self.buka_popup_pilihan_produk(match_manual)
            self.ent_search.delete(0, tk.END)
        else:
            messagebox.showwarning("Tidak Ditemukan", f"Produk dengan kata kunci '{raw_input}' tidak ditemukan!")

    def buka_popup_pilihan_produk(self, df_pilihan):
        popup = tk.Toplevel(self.root)
        popup.title("Pilih Produk")
        self.center_popup(popup, 580, 380)
        popup.transient(self.root)
        popup.grab_set()
        popup.bind("<Escape>", lambda e: (popup.destroy(), self.ent_search.focus_force()))

        ttk.Label(popup, text="Gunakan Panah Atas/Bawah lalu tekan ENTER untuk memilih:", font=("Arial", 9, "bold")).pack(pady=8)
        frame_list = ttk.Frame(popup)
        frame_list.pack(fill="both", expand=True, padx=10, pady=5)

        cols = ("barcode", "judul", "harga")
        tree_select = ttk.Treeview(frame_list, columns=cols, show="headings", selectmode="browse")
        tree_select.heading("barcode", text="Barcode")
        tree_select.heading("judul", text="Judul")
        tree_select.heading("harga", text="Harga")
        tree_select.column("barcode", width=120)
        tree_select.column("judul", width=310)
        tree_select.column("harga", width=100, anchor="e")

        sb = ttk.Scrollbar(frame_list, orient="vertical", command=tree_select.yview)
        tree_select.configure(yscroll=sb.set)
        sb.pack(side="right", fill="y")
        tree_select.pack(side="left", fill="both", expand=True)

        for _, row in df_pilihan.iterrows():
            harga_val = float(row['harga'])
            tree_select.insert("", "end", values=(row['barcode'], row['judul'], f"Rp. {harga_val:,.0f}"))

        children = tree_select.get_children()
        if children:
            tree_select.selection_set(children[0])
            tree_select.focus_set()
            tree_select.focus(children[0])

        def pilih_item_terpilih(event=None):
            selected = tree_select.selection()
            if not selected:
                return
            idx = tree_select.index(selected[0])
            item_selected = df_pilihan.iloc[idx]
            popup.destroy()
            self.buka_popup_input_produk(item_selected)

        tree_select.bind("<Return>", pilih_item_terpilih)
        popup.bind("<Return>", pilih_item_terpilih)
        ttk.Button(popup, text="Pilih Produk Ini (ENTER)", command=pilih_item_terpilih).pack(pady=8)

    def buka_popup_input_produk(self, item):
        popup = tk.Toplevel(self.root)
        popup.title("Input Jumlah & Diskon")
        self.center_popup(popup, 400, 260)
        popup.transient(self.root)
        popup.grab_set()
        popup.bind("<Escape>", lambda e: (popup.destroy(), self.ent_search.focus_force()))

        ttk.Label(popup, text=f"Produk: {item['judul']}", font=("Arial", 10, "bold"), wraplength=380).pack(pady=(10, 2))
        harga_norm = float(item['harga'])
        ttk.Label(popup, text=f"Harga Normal: Rp. {harga_norm:,.0f}", foreground="gray").pack(pady=(0, 10))

        row_qty = ttk.Frame(popup)
        row_qty.pack(pady=5)
        ttk.Label(row_qty, text="Jumlah (Qty):").pack(side="left", padx=5)
        ent_qty = ttk.Entry(row_qty, width=10, font=("Arial", 10, "bold"))
        ent_qty.insert(0, "1")
        ent_qty.pack(side="left", padx=5)
        ent_qty.focus_force()
        ent_qty.selection_range(0, tk.END)

        frame_diskon = ttk.LabelFrame(popup, text=" Opsi Diskon ")
        frame_diskon.pack(fill="x", padx=20, pady=5)
        diskon_var = tk.StringVar(value="0")

        row_radio = ttk.Frame(frame_diskon)
        row_radio.pack(pady=5)
        ttk.Radiobutton(row_radio, text="0%", variable=diskon_var, value="0").pack(side="left", padx=5)
        ttk.Radiobutton(row_radio, text="5%", variable=diskon_var, value="5").pack(side="left", padx=5)
        ttk.Radiobutton(row_radio, text="20%", variable=diskon_var, value="20").pack(side="left", padx=5)
        ttk.Radiobutton(row_radio, text="25%", variable=diskon_var, value="25").pack(side="left", padx=5)

        row_custom = ttk.Frame(frame_diskon)
        row_custom.pack(pady=5)
        ttk.Label(row_custom, text="Diskon Custom (%):").pack(side="left", padx=5)
        ent_custom_diskon = ttk.Entry(row_custom, width=8)
        ent_custom_diskon.insert(0, "0")
        ent_custom_diskon.pack(side="left", padx=5)

        def submit_item(event=None):
            try:
                added_qty = int(ent_qty.get().strip())
                if added_qty <= 0: raise ValueError()
            except:
                messagebox.showwarning("Peringatan", "Jumlah Qty harus angka bulat positif!", parent=popup)
                return

            custom_val = ent_custom_diskon.get().strip()
            if custom_val != "0" and custom_val != "":
                try: diskon_percent = float(custom_val)
                except: diskon_percent = float(diskon_var.get())
            else:
                diskon_percent = float(diskon_var.get())

            harga_akhir = harga_norm * (1 - (diskon_percent / 100))
            barcode_str = str(item['barcode'])

            existing_item = None
            for cart_item in self.cart:
                if cart_item['barcode'] == barcode_str and cart_item['diskon'] == diskon_percent:
                    existing_item = cart_item
                    break

            if existing_item:
                existing_item['jumlah'] += added_qty
                existing_item['total'] = existing_item['jumlah'] * harga_akhir
            else:
                total = harga_akhir * added_qty
                self.cart.append({
                    "barcode": barcode_str, "judul": str(item['judul']), "jumlah": added_qty,
                    "harga_normal": harga_norm, "diskon": diskon_percent, "harga_akhir": harga_akhir, "total": total
                })

            self.update_tabel_keranjang()
            popup.destroy()
            self.ent_search.focus_force()

        popup.bind("<Return>", submit_item)
        btn_row = ttk.Frame(popup)
        btn_row.pack(pady=10)
        ttk.Button(btn_row, text="OK (ENTER)", command=submit_item).pack(side="left", padx=5)
        ttk.Button(btn_row, text="Batal (ESC)", command=lambda: (popup.destroy(), self.ent_search.focus_force())).pack(side="left", padx=5)

    def update_tabel_keranjang(self):
        for row in self.tree.get_children():
            self.tree.delete(row)
        grand_total = 0
        total_qty = 0

        for idx, item in enumerate(self.cart, 1):
            diskon_str = f"{int(item['diskon'])}%" if float(item['diskon']).is_integer() else f"{item['diskon']}%"
            self.tree.insert("", "end", values=(
                idx, item['barcode'], item['judul'], item['jumlah'],
                f"Rp. {item['harga_normal']:,.0f}", diskon_str,
                f"Rp. {item['harga_akhir']:,.0f}", f"Rp. {item['total']:,.0f}"
            ))
            grand_total += item['total']
            total_qty += item['jumlah']

        self.lbl_total.config(text=f"TOTAL BELANJA: Rp. {grand_total:,.0f}")
        self.lbl_total_qty.config(text=f"TOTAL ITEM / QTY: {total_qty:,} Pcs")

    def reset_transaksi(self):
        if messagebox.askyesno("Konfirmasi", "Bersihkan keranjang belanja saat ini?"):
            self.cart.clear()
            self.update_tabel_keranjang()
            self.ent_search.focus_force()

    def cetak_nota_thermal_80mm_custom(self, no_trx, metode, bayar, kembali, catatan_member, waktu_str=None, items_source=None, is_reprint=False):
        items_to_print = items_source if items_source is not None else self.cart
        grand_total = sum(item['total'] for item in items_to_print)
        if not waktu_str:
            waktu_str = datetime.now().strftime("%d/%m/%Y %H:%M:%S")
        
        lebar = self.config_nota["lebar_kertas"]
        line_sep = "-" * lebar
        
        # Gunakan nama kasir dari config, fallback ke "Admin"
        nama_kasir = self.config_nota.get("nama_kasir", "Admin")
        
        struk = []
        struk.append("".center(lebar))
        if self.config_nota["tampilkan_logo_teks"]:
            struk.append(self.config_nota["nama_toko"].center(lebar))
            if self.config_nota["sub_nama"]:
                struk.append(self.config_nota["sub_nama"].center(lebar))
                
        struk.append(self.config_nota["alamat_toko"].center(lebar))
        if self.config_nota["tampilkan_telp"] and self.config_nota["telepon_toko"]:
            struk.append(self.config_nota["telepon_toko"].center(lebar))
            
        struk.append(line_sep)
        if is_reprint:
            struk.append("*** RE-PRINT NOTA KASIR ***".center(lebar))
        struk.append(f"No. Trx  : #{no_trx}".ljust(lebar // 2) + f"Kasir : {nama_kasir}".rjust(lebar - (lebar // 2)))
        struk.append(f"Waktu   : {waktu_str}".ljust(lebar))
        if catatan_member:
            struk.append(f"Member  : {catatan_member}".ljust(lebar))
        struk.append(line_sep)
        
        for item in items_to_print:
            judul = item['judul'][:lebar]
            struk.append(judul)
            qty_price = f"  {item['jumlah']} x @{item['harga_akhir']:,.0f}"
            subtotal = f"Rp. {item['total']:,.0f}"
            p_lebar_sub = len(qty_price)
            spasi_tengah = max(2, lebar - p_lebar_sub - len(subtotal))
            struk.append(qty_price + (" " * spasi_tengah) + subtotal)
            if item.get('diskon', 0) > 0:
                struk.append(f"  (Diskon {item['diskon']}%)".ljust(lebar))
                
        struk.append(line_sep)
        struk.append(f"TOTAL BELANJA".ljust(lebar // 2) + f"Rp. {grand_total:,.0f}".rjust(lebar - (lebar // 2)))
        struk.append(f"PEMBAYARAN ({metode})".ljust(lebar // 2) + f"Rp. {bayar:,.0f}".rjust(lebar - (lebar // 2)))
        if metode == "TUNAI":
            struk.append(f"KEMBALIAN".ljust(lebar // 2) + f"Rp. {kembali:,.0f}".rjust(lebar - (lebar // 2)))
        struk.append(line_sep)
        
        if self.config_nota["tampilkan_footer"]:
            if self.config_nota["info_tambahan"]:
                struk.append(self.config_nota["info_tambahan"].center(lebar))
            if self.config_nota["pesan_penutup"]:
                struk.append(self.config_nota["pesan_penutup"].center(lebar))
            
        # Perintah Paper Cut ESC/POS untuk pemotong otomatis
        PAPER_CUT = "\x1dV\x41\x00"
        text_struk_final = "\n".join(struk) + "\n\n\n\n" + PAPER_CUT
        
        # PROSES CETAK RAW KE PRINTER
        try:
            if os.name == 'nt':
                try:
                    import win32print
                except ImportError:
                    messagebox.showerror("Error Printer", "Modul 'pywin32' tidak ditemukan.\nSilakan instal via: pip install pywin32")
                    return
                
                # Ambil printer default pada sistem Windows
                printer_name = win32print.GetDefaultPrinter()
                
                # Buka koneksi printer & kirim data RAW secara langsung
                hPrinter = win32print.OpenPrinter(printer_name)
                try:
                    hJob = win32print.StartDocPrinter(hPrinter, 1, ("Nota Transaksi Kasir", None, "RAW"))
                    win32print.StartPagePrinter(hPrinter)
                    win32print.WritePrinter(hPrinter, text_struk_final.encode('utf-8', errors='ignore'))
                    win32print.EndPagePrinter(hPrinter)
                    win32print.EndDocPrinter(hPrinter)
                finally:
                    win32print.ClosePrinter(hPrinter)
            else:
                # Untuk OS Linux / Mac OS
                with tempfile.NamedTemporaryFile(delete=False, suffix=".txt", mode="w", encoding="utf-8") as tmp:
                    tmp.write(text_struk_final)
                    tmp_path = tmp.name
                os.system(f"lpr -o raw {tmp_path}")
                os.remove(tmp_path)
                
        except Exception as e:
            print("Gagal mencetak otomatis:", e)
            messagebox.showerror("Error Printer", f"Gagal mencetak struk ke printer thermal:\n{str(e)}")

    def buka_popup_pembayaran(self):
        if not self.cart:
            messagebox.showwarning("Peringatan", "Keranjang belanja masih kosong!")
            return
        if not self.target_file_path or not self.combo_sheet.get():
            messagebox.showwarning("Peringatan", "Silakan pilih File Target Output & Sheet Target terlebih dahulu!")
            return

        grand_total = sum(item['total'] for item in self.cart)
        popup = tk.Toplevel(self.root)
        popup.title("Pembayaran")
        self.center_popup(popup, 380, 440)
        popup.transient(self.root)
        popup.grab_set()
        popup.bind("<Escape>", lambda e: (popup.destroy(), self.ent_search.focus_force()))

        ttk.Label(popup, text="Total Transaksi:", font=("Arial", 11, "bold")).pack(pady=(10, 2))
        ttk.Label(popup, text=f"Rp. {grand_total:,.0f}", font=("Arial", 16, "bold"), foreground="blue").pack(pady=(0, 5))

        ada_diskon = any(item['diskon'] > 0 for item in self.cart)
        frame_member_pay = ttk.LabelFrame(popup, text=" Catatan Member ")
        frame_member_pay.pack(fill="x", padx=30, pady=5)
        ent_catatan_pay = ttk.Entry(frame_member_pay, font=("Arial", 10))
        ent_catatan_pay.pack(fill="x", padx=10, pady=5)

        if ada_diskon:
            ent_catatan_pay.config(state="normal")
        else:
            ent_catatan_pay.config(state="disabled")
            ent_catatan_pay.insert(0, "(Tidak ada diskon / Non-Member)")

        ttk.Label(popup, text="Pilih Jenis Pembayaran:").pack(anchor="w", padx=30, pady=(5, 0))
        pay_var = tk.StringVar(value="TUNAI")
        frame_tunai_input = ttk.Frame(popup)
        
        ttk.Label(frame_tunai_input, text="Nominal Bayar (Rp):").pack(anchor="w", pady=(5, 2))
        ent_bayar = ttk.Entry(frame_tunai_input, font=("Arial", 11, "bold"))
        ent_bayar.pack(fill="x")
        lbl_kembalian = ttk.Label(frame_tunai_input, text="Kembalian: Rp. 0", font=("Arial", 10, "bold"), foreground="green")
        lbl_kembalian.pack(anchor="w", pady=(5, 0))

        def hitung_kembalian(event=None):
            raw = ent_bayar.get().replace(".", "").replace(",", "").strip()
            if raw.isdigit():
                bayar = float(raw)
                kembali = bayar - grand_total
                if kembali >= 0:
                    lbl_kembalian.config(text=f"Kembalian: Rp. {kembali:,.0f}", foreground="green")
                else:
                    lbl_kembalian.config(text=f"Kurang: Rp. {abs(kembali):,.0f}", foreground="red")
            else:
                lbl_kembalian.config(text="Kembalian: Rp. 0", foreground="green")

        ent_bayar.bind("<KeyRelease>", hitung_kembalian)

        def toggle_pembayaran():
            if pay_var.get() == "TUNAI":
                frame_tunai_input.pack(fill="x", padx=30, pady=5)
                ent_bayar.focus_force()
            else:
                frame_tunai_input.pack_forget()

        row_radio = ttk.Frame(popup)
        row_radio.pack(fill="x", padx=30, pady=5)
        ttk.Radiobutton(row_radio, text="TUNAI", variable=pay_var, value="TUNAI", command=toggle_pembayaran).pack(side="left", padx=10)
        ttk.Radiobutton(row_radio, text="NON TUNAI", variable=pay_var, value="NON TUNAI", command=toggle_pembayaran).pack(side="left", padx=10)
        toggle_pembayaran()

        def simpan_dan_proses(event=None):
            catatan_val = ""
            if ada_diskon:
                catatan_val = ent_catatan_pay.get().strip()
                if not catatan_val:
                    messagebox.showwarning("Peringatan", "Catatan Member wajib diisi karena ada diskon!", parent=popup)
                    ent_catatan_pay.focus_force()
                    return

            metode = pay_var.get()
            tunai_val = 0
            nontunai_val = 0
            bayar_val = grand_total
            kembali_val = 0

            if metode == "TUNAI":
                raw_bayar = ent_bayar.get().replace(".", "").replace(",", "").strip()
                if not raw_bayar.isdigit() or float(raw_bayar) < grand_total:
                    messagebox.showwarning("Peringatan", "Nominal pembayaran Tunai kurang!", parent=popup)
                    return
                bayar_val = float(raw_bayar)
                tunai_val = grand_total
                kembali_val = bayar_val - grand_total
            else:
                nontunai_val = grand_total

            try:
                sheet_name = self.combo_sheet.get()
                ws = self.wb_target[sheet_name]
                self.cek_dan_buat_header_excel(ws)
                self.hitung_nomor_transaksi_berikutnya()

                no_trx_num = self.no_trx_counter
                timestamp_str = datetime.now().strftime("%H:%M:%S")
                currency_format = '"Rp" #,##0'
                
                thin_border = Border(
                    left=Side(style='thin', color='000000'), right=Side(style='thin', color='000000'),
                    top=Side(style='thin', color='000000'), bottom=Side(style='thin', color='000000')
                )

                all_used_cols = ["A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K"]

                for idx_item, item in enumerate(self.cart):
                    max_r = ws.max_row + 1
                    if idx_item == 0: ws[f"A{max_r}"] = no_trx_num

                    ws[f"B{max_r}"] = item['barcode']
                    ws[f"C{max_r}"] = item['judul']
                    ws[f"D{max_r}"] = item['jumlah']
                    
                    ws[f"E{max_r}"] = item['harga_akhir']
                    ws[f"E{max_r}"].number_format = currency_format

                    ws[f"F{max_r}"] = item['total']
                    ws[f"F{max_r}"].number_format = currency_format

                    if idx_item == 0:
                        if metode == "TUNAI":
                            ws[f"G{max_r}"] = tunai_val
                            ws[f"G{max_r}"].number_format = currency_format
                        else:
                            ws[f"H{max_r}"] = nontunai_val
                            ws[f"H{max_r}"].number_format = currency_format

                    ws[f"I{max_r}"] = item['diskon']
                    ws[f"J{max_r}"] = catatan_val if ada_diskon else ""
                    if idx_item == 0: ws[f"K{max_r}"] = timestamp_str

                    for col_l in all_used_cols:
                        ws[f"{col_l}{max_r}"].border = thin_border

                self.wb_target.save(self.target_file_path)
                
                # Invalidasi cache laporan karena ada data baru
                self.cached_report_df = None
                self.cached_sheet_name = None
                
                self.cetak_nota_thermal_80mm_custom(no_trx_num, metode, bayar_val, kembali_val, catatan_val if ada_diskon else "", is_reprint=False)

                messagebox.showinfo("Sukses", f"Transaksi No. {no_trx_num} berhasil disimpan & dicetak!")
                
                self.cart.clear()
                self.update_tabel_keranjang()
                self.hitung_nomor_transaksi_berikutnya()
                popup.destroy()
                self.ent_search.focus_force()
                self.refresh_data_laporan()
            except Exception as e:
                messagebox.showerror("Error Simpan", f"Gagal menyimpan transaksi:\n{str(e)}", parent=popup)

        popup.bind("<Return>", simpan_dan_proses)
        ttk.Button(popup, text="SIMPAN & CETAK NOTA (ENTER)", command=simpan_dan_proses).pack(side="bottom", pady=10)

if __name__ == "__main__":
    root = tk.Tk()
    app = KasirApp(root)
    root.mainloop()
